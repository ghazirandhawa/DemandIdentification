from __future__ import annotations

import os
from pathlib import Path
from typing import Any

_DEFAULT_MAX_CHARS_SCOPE = 20_000
_DEFAULT_MAX_CHARS_ICP = 55_000
_DEFAULT_MAX_CHARS_DECK = 45_000
_DEFAULT_MAX_TOTAL = 120_000


def _clip(text: str, max_chars: int) -> str:
    t = (text or "").strip()
    if len(t) <= max_chars:
        return t
    return t[: max_chars - 80] + "\n\n[... truncated for prompt size ...]\n"


def _load_txt(path: Path) -> str | None:
    if not path.is_file():
        return None
    return path.read_text(encoding="utf-8", errors="replace")


def _docx_table_text(doc: Any) -> list[str]:
    rows_out: list[str] = []
    try:
        for table in doc.tables:
            for row in table.rows:
                cells = [(c.text or "").strip() for c in row.cells]
                cells = [c for c in cells if c]
                if cells:
                    rows_out.append(" | ".join(cells))
    except Exception:
        pass
    return rows_out


def _load_docx(path: Path) -> str | None:
    if not path.is_file():
        return None
    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError:
        return None
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        s = (p.text or "").strip()
        if s:
            parts.append(s)
    tbl = _docx_table_text(doc)
    if tbl:
        parts.append("=== TABLES (row text) ===")
        parts.extend(tbl)
    return "\n".join(parts) if parts else ""


def _pptx_collect_shape_text(shape: Any, out: list[str]) -> None:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore[import-untyped]
    except ImportError:
        if hasattr(shape, "text"):
            t = (getattr(shape, "text", None) or "").strip()
            if t:
                out.append(t)
        return
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in getattr(shape, "shapes", []):
                _pptx_collect_shape_text(child, out)
            return
    except Exception:
        pass
    if hasattr(shape, "text"):
        t = (getattr(shape, "text", None) or "").strip()
        if t:
            out.append(t)


def _load_pptx(path: Path) -> str | None:
    if not path.is_file():
        return None
    try:
        from pptx import Presentation  # type: ignore[import-untyped]
    except ImportError:
        return None
    prs = Presentation(str(path))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            _pptx_collect_shape_text(shape, texts)
    return "\n".join(texts) if texts else ""


def datapilot_docs_dir(repo_root: Path | None = None) -> Path:
    root = repo_root or Path(__file__).resolve().parents[2]
    return root / "datapilot_docs"


def _iter_datapilot_files(d: Path) -> list[Path]:
    """
    List regular files under ``datapilot_docs`` (runtime filesystem).

    Skips WSL/Windows junk like ``*:Zone.Identifier`` so discovery stays stable.
    """
    if not d.is_dir():
        return []
    out: list[Path] = []
    for p in d.iterdir():
        if not p.is_file():
            continue
        name = p.name
        if ":" in name or name.endswith("Zone.Identifier"):
            continue
        out.append(p)
    out.sort(key=lambda x: x.name.lower())
    return out


def _pick_icp_docx(d: Path) -> Path | None:
    paths = _iter_datapilot_files(d)
    docxs = [p for p in paths if p.suffix.lower() == ".docx"]
    if not docxs:
        return None
    preferred = [
        "DataPilot_ICP_v3.docx",
        (os.environ.get("DATAPILOT_ICP_DOCX") or "").strip(),
    ]
    for pref in preferred:
        if not pref:
            continue
        for p in docxs:
            if p.name == pref:
                return p
    for p in docxs:
        if "icp" in p.stem.lower():
            return p
    return docxs[0] if len(docxs) == 1 else None


def _pick_sales_deck_pptx(d: Path) -> Path | None:
    paths = _iter_datapilot_files(d)
    pptxs = [p for p in paths if p.suffix.lower() == ".pptx"]
    if not pptxs:
        return None
    preferred = (
        "Data Pilot Sales Deck Complete.pptx",
        "Data_Pilot_Sales_Deck_Complete.pptx",
    )
    for pref in preferred:
        for p in pptxs:
            if p.name == pref:
                return p
    for p in pptxs:
        stem = p.stem.lower()
        if "sales" in stem and "deck" in stem:
            return p
    for p in pptxs:
        stem = p.stem.lower()
        if "deck" in stem or "sales" in stem:
            return p
    return pptxs[0] if len(pptxs) == 1 else None


def load_datapilot_bundle(
    *,
    repo_root: Path | None = None,
    max_chars_per_source: int | None = None,
    max_chars_scope: int | None = None,
    max_chars_icp: int | None = None,
    max_chars_deck: int | None = None,
    max_total_chars: int | None = None,
) -> dict[str, Any]:
    """
    Load ICP docx, project scope txt, and sales deck pptx when present.

    Missing files or optional deps (python-docx / python-pptx) are reported in
    ``meta``; the pipeline should still run with whatever context is available.

    If ``max_chars_per_source`` is set, it overrides the per-section defaults (legacy).
    Otherwise use ``max_chars_scope`` / ``max_chars_icp`` / ``max_chars_deck`` when provided.
    The concatenated bundle is clipped to ``max_total_chars`` (default ``_DEFAULT_MAX_TOTAL``).
    """
    if max_chars_per_source is not None:
        cap_scope = cap_icp = cap_deck = int(max_chars_per_source)
    else:
        cap_scope = int(max_chars_scope if max_chars_scope is not None else _DEFAULT_MAX_CHARS_SCOPE)
        cap_icp = int(max_chars_icp if max_chars_icp is not None else _DEFAULT_MAX_CHARS_ICP)
        cap_deck = int(max_chars_deck if max_chars_deck is not None else _DEFAULT_MAX_CHARS_DECK)
    cap_total = int(max_total_chars if max_total_chars is not None else _DEFAULT_MAX_TOTAL)

    d = datapilot_docs_dir(repo_root)
    meta: dict[str, str] = {}
    parts: list[str] = []

    scope_path = d / "project_scope.txt"
    scope = _load_txt(scope_path)
    if scope:
        parts.append("=== DataPilot project_scope.txt ===\n" + _clip(scope, cap_scope))
        meta["project_scope"] = str(scope_path)
    else:
        meta["project_scope"] = f"missing:{scope_path}"

    icp_path = _pick_icp_docx(d)
    if icp_path:
        icp_text = _load_docx(icp_path)
        if icp_text is None:
            meta["icp"] = f"unreadable_or_no_docx_lib:{icp_path}"
        elif not icp_text.strip():
            meta["icp"] = f"empty:{icp_path}"
        else:
            parts.append("=== DataPilot ICP (docx) ===\n" + _clip(icp_text, cap_icp))
            meta["icp"] = str(icp_path)
    else:
        meta["icp"] = f"no_docx_found_under:{d}"

    deck_path = _pick_sales_deck_pptx(d)
    if deck_path:
        deck_text = _load_pptx(deck_path)
        if deck_text is None:
            meta["sales_deck"] = f"unreadable_or_no_pptx_lib:{deck_path}"
        elif not deck_text.strip():
            meta["sales_deck"] = f"empty:{deck_path}"
        else:
            parts.append("=== DataPilot sales deck (pptx text) ===\n" + _clip(deck_text, cap_deck))
            meta["sales_deck"] = str(deck_path)
    else:
        meta["sales_deck"] = f"no_pptx_found_under:{d}"

    combined = "\n\n".join(parts).strip()
    combined = _clip(combined, cap_total)
    meta["context_limits"] = (
        f"scope<={cap_scope}, icp<={cap_icp}, deck<={cap_deck}, total<={cap_total}"
    )
    return {"text": combined, "meta": meta}
